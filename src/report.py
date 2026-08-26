"""
Writes the checker's results into a real, formatted Excel report --
the "Erstellung, Optimierung und Weiterentwicklung ... von Reports und
Kennzahlen" ask from the posting -- with one sheet per issue type plus
a summary sheet with headline counts.
"""

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_df_sheet(wb, title, df):
    ws = wb.create_sheet(title)
    if df.empty:
        ws.append(["Keine Auffälligkeiten gefunden"])
        return
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for _, row in df.iterrows():
        ws.append(list(row))
    for col_cells in ws.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 40)


def write_report(results, path="output/milestone_check_report.xlsx"):
    wb = Workbook()
    summary = wb.active
    summary.title = "Zusammenfassung"
    summary.append(["Kennzahl", "Wert"])
    for cell in summary[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    summary.append(["Geprüfte Standorte", results["total_sites"]])
    summary.append(["Gefundene Auffälligkeiten gesamt", results["total_issues"]])
    summary.append(["Falsche Meilenstein-Reihenfolge", len(results["out_of_order"])])
    summary.append(["SLA-/Anforderungsverletzungen", len(results["sla_breaches"])])
    summary.append(["Fehlende Meilenstein-Daten", len(results["missing_dates"])])
    summary.append(["Doppelte Standort-Einträge", len(results["duplicates"])])
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 14

    _write_df_sheet(wb, "Falsche Reihenfolge", results["out_of_order"])
    _write_df_sheet(wb, "SLA-Verletzungen", results["sla_breaches"])
    _write_df_sheet(wb, "Fehlende Daten", results["missing_dates"])
    _write_df_sheet(wb, "Duplikate", results["duplicates"])

    wb.save(path)
    return path
