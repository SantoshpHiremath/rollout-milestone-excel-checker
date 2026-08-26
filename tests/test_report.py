import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from generate_data import generate_sites, write_sites_workbook, write_requirements_workbook
from checker import load_sites, load_requirements, run_all_checks
from report import write_report


@pytest.fixture(scope="module")
def report_path(tmp_path_factory):
    rows, _ = generate_sites(n=300, seed=3)
    out_dir = tmp_path_factory.mktemp("data")
    sites_path = str(out_dir / "sites.xlsx")
    reqs_path = str(out_dir / "requirements.xlsx")
    write_sites_workbook(rows, sites_path)
    write_requirements_workbook(reqs_path)

    sites_df = load_sites(sites_path)
    requirements_df = load_requirements(reqs_path)
    results = run_all_checks(sites_df, requirements_df)

    report_dir = tmp_path_factory.mktemp("output")
    path = str(report_dir / "report.xlsx")
    write_report(results, path)
    return path, results


def test_write_report_creates_a_real_xlsx_file(report_path):
    path, _ = report_path
    assert os.path.exists(path)
    wb = openpyxl.load_workbook(path)
    assert wb is not None


def test_write_report_has_expected_sheets(report_path):
    path, _ = report_path
    wb = openpyxl.load_workbook(path)
    expected = {"Zusammenfassung", "Falsche Reihenfolge", "SLA-Verletzungen", "Fehlende Daten", "Duplikate"}
    assert expected <= set(wb.sheetnames)


def test_write_report_summary_sheet_total_matches_results(report_path):
    path, results = report_path
    wb = openpyxl.load_workbook(path)
    ws = wb["Zusammenfassung"]
    rows = list(ws.iter_rows(values_only=True))
    summary = {row[0]: row[1] for row in rows[1:]}
    assert summary["Geprüfte Standorte"] == results["total_sites"]
    assert summary["Gefundene Auffälligkeiten gesamt"] == results["total_issues"]


def test_write_report_sla_sheet_row_count_matches_results(report_path):
    path, results = report_path
    wb = openpyxl.load_workbook(path)
    ws = wb["SLA-Verletzungen"]
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    if len(results["sla_breaches"]) == 0:
        assert len(data_rows) == 1  # "Keine Auffälligkeiten gefunden" placeholder row
    else:
        assert len(data_rows) == len(results["sla_breaches"])
